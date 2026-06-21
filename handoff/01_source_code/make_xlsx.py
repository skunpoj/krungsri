from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import random
random.seed(7)

wb = Workbook()
ws = wb.active
ws.title = "ข้อมูลส่งออก"

FNT = "Arial"
hdr_fill = PatternFill("solid", start_color="1F4E5F")
hdr_font = Font(name=FNT, bold=True, color="FFFFFF", size=10)
cell_font = Font(name=FNT, size=10)
thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center")
right = Alignment(horizontal="right")

headers = ["เดือน","สินค้า","HS Code","ประเทศปลายทาง","ภูมิภาค","จำนวน (กก.)","มูลค่า (USD)","Incoterms","ผู้ซื้อ"]
for c,h in enumerate(headers,1):
    cell = ws.cell(1,c,h); cell.fill=hdr_fill; cell.font=hdr_font; cell.alignment=center; cell.border=border

products = [
    ("ข้าวหอมมะลิ 100%","1006.30",1.20),
    ("แป้งมันสำปะหลัง","1108.14",0.50),
    ("สับปะรดกระป๋อง","2008.20",0.95),
    ("ถุงมือยางทางการแพทย์","4015.12",3.40),
    ("ชิ้นส่วนยานยนต์","8708.99",6.80),
]
# country: (name, region, base demand weight, monthly growth factor)
countries = [
    ("สหรัฐอเมริกา","อเมริกาเหนือ",1.0,0.985),   # declining (tariffs)
    ("จีน","เอเชีย",0.9,1.03),                    # growing
    ("ญี่ปุ่น","เอเชีย",0.8,1.005),
    ("เยอรมนี","ยุโรป",0.7,1.0),
    ("เนเธอร์แลนด์","ยุโรป",0.5,1.008),
    ("อินเดีย","เอเชีย",0.45,1.05),               # fast growing
    ("เวียดนาม","อาเซียน",0.6,1.02),
    ("สหรัฐอาหรับเอมิเรตส์","ตะวันออกกลาง",0.4,1.025),
    ("ออสเตรเลีย","โอเชียเนีย",0.5,1.004),
    ("เกาหลีใต้","เอเชีย",0.55,1.01),
]
buyers = {
 "สหรัฐอเมริกา":"Pacific Foods Inc.","จีน":"Shanghai Trade Co.","ญี่ปุ่น":"Tokyo Import KK",
 "เยอรมนี":"EuroFood GmbH","เนเธอร์แลนด์":"Rotterdam B.V.","อินเดีย":"Mumbai Traders Pvt",
 "เวียดนาม":"Saigon Distributors","สหรัฐอาหรับเอมิเรตส์":"Gulf Source LLC",
 "ออสเตรเลีย":"Sydney Wholesale Pty","เกาหลีใต้":"Seoul Mart Co.",
}
months = [f"2025-{m:02d}" for m in range(1,13)] + [f"2026-{m:02d}" for m in range(1,6)]

rows=[]
for mi,mon in enumerate(months):
    # pick a handful of shipments per month
    picks = random.sample(countries, k=random.randint(3,5))
    for (cn,reg,w,g) in picks:
        prod,hs,price = random.choice(products)
        growth = g**mi
        qty = int(random.randint(8000,42000)*w*growth/1000)*1000
        if qty < 1000: qty = 1000
        val = round(qty*price*random.uniform(0.95,1.08))
        inco = random.choice(["FOB Bangkok","CIF","FOB Laem Chabang"])
        rows.append([mon,prod,hs,cn,reg,qty,val,inco,buyers[cn]])

r=2
for row in rows:
    for c,v in enumerate(row,1):
        cell=ws.cell(r,c,v); cell.font=cell_font; cell.border=border
        if c in (6,7): cell.alignment=right; cell.number_format="#,##0"
        if c in (1,3): cell.alignment=center
    r+=1
last=r-1

# column widths
for col,wd in zip("ABCDEFGHI",[10,22,10,20,14,13,14,16,20]):
    ws.column_dimensions[col].width=wd
ws.freeze_panes="A2"

# total row
ws.cell(r,5,"รวมทั้งหมด").font=Font(name=FNT,bold=True)
ws.cell(r,6,f"=SUM(F2:F{last})").font=Font(name=FNT,bold=True); ws.cell(r,6).number_format="#,##0"; ws.cell(r,6).alignment=right
ws.cell(r,7,f"=SUM(G2:G{last})").font=Font(name=FNT,bold=True); ws.cell(r,7).number_format="#,##0"; ws.cell(r,7).alignment=right

# Summary sheet
ws2=wb.create_sheet("สรุปตามตลาด")
sh=["ประเทศปลายทาง","มูลค่ารวม (USD)","จำนวนชิปเมนต์","สัดส่วน %"]
for c,h in enumerate(sh,1):
    cell=ws2.cell(1,c,h); cell.fill=hdr_fill; cell.font=hdr_font; cell.alignment=center; cell.border=border
cn_list=[c[0] for c in countries]
for i,cn in enumerate(cn_list,2):
    ws2.cell(i,1,cn).font=cell_font; ws2.cell(i,1).border=border
    ws2.cell(i,2,f"=SUMIF('ข้อมูลส่งออก'!$D$2:$D${last},A{i},'ข้อมูลส่งออก'!$G$2:$G${last})")
    ws2.cell(i,3,f"=COUNTIF('ข้อมูลส่งออก'!$D$2:$D${last},A{i})")
    ws2.cell(i,2).number_format="#,##0"; ws2.cell(i,2).alignment=right; ws2.cell(i,2).font=cell_font; ws2.cell(i,2).border=border
    ws2.cell(i,3).alignment=center; ws2.cell(i,3).font=cell_font; ws2.cell(i,3).border=border
tot=len(cn_list)+2
ws2.cell(tot,1,"รวม").font=Font(name=FNT,bold=True)
ws2.cell(tot,2,f"=SUM(B2:B{tot-1})").font=Font(name=FNT,bold=True); ws2.cell(tot,2).number_format="#,##0"; ws2.cell(tot,2).alignment=right
ws2.cell(tot,3,f"=SUM(C2:C{tot-1})").font=Font(name=FNT,bold=True); ws2.cell(tot,3).alignment=center
for i in range(2,tot):
    ws2.cell(i,4,f"=B{i}/$B${tot}"); ws2.cell(i,4).number_format="0.0%"; ws2.cell(i,4).alignment=right; ws2.cell(i,4).font=cell_font; ws2.cell(i,4).border=border
for col,wd in zip("ABCD",[22,18,16,12]):
    ws2.column_dimensions[col].width=wd

# Notes sheet
ws3=wb.create_sheet("คำอธิบาย")
notes=[
 ["ไฟล์ตัวอย่างข้อมูลส่งออก (สำหรับเดโม AI หาตลาด)",""],
 ["",""],
 ["วิธีใช้ในเวิร์กช็อป:",""],
 ["1) อัปโหลดไฟล์นี้เข้า ChatGPT / Claude (เปิด Web Search)",""],
 ["2) พิมพ์ Keyword เช่น:",""],
 ["   \"จากไฟล์ข้อมูลส่งออกนี้ ช่วยวิเคราะห์ Top 5 ตลาดของฉัน และตลาดที่กำลังเติบโตปี 2026\"",""],
 ["3) ลองเปลี่ยน Keyword เพื่อดูผลอีกแบบ เช่น ดูคู่แข่ง / ความเสี่ยงภาษี / สินค้าที่ควรโฟกัส",""],
 ["",""],
 ["หมายเหตุ: ข้อมูลนี้เป็นตัวอย่างสมมติ ไม่ใช่ข้อมูลจริง — เปลี่ยนเป็นไฟล์ของธุรกิจคุณได้",""],
]
for i,(a,b) in enumerate(notes,1):
    ws3.cell(i,1,a).font=Font(name=FNT, bold=(i==1), size=12 if i==1 else 10)
ws3.column_dimensions["A"].width=90

wb.save("/home/claude/Sample_Thai_Export_Data.xlsx")
print("saved; data rows:",len(rows))
